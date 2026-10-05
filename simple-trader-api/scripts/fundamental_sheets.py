"""
Per-symbol fundamental sheets for the daily watchlist workbook.

Each symbol gets a sheet named after it: verdict banner, key figures, a 3-4 line summary, then every
rule as Check | Value | Result (like a worked example). The Watchlist sheet's symbol cell links to it
and the sheet links back.
"""
import re

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.hyperlink import Hyperlink

from scripts.fundamental_rules import GATE_NAMES, HARD_GATES, summary_notes

VERDICT_FILLS = {"PASS": "C6EFCE", "WATCH": "FFEB9C", "REJECT": "FFC7CE", "FINANCIAL": "DDEBF7", "NO DATA": "D9D9D9"}
RESULT_STYLE = {  # level -> (label, fill)
    "pass": ("Pass", "C6EFCE"), "bonus": ("Bonus", "C6EFCE"), "ok": ("Neutral", "F2F2F2"),
    "warn": ("Warn", "FFEB9C"), "fail": ("Fail", "F8CBAD"), "info": ("Info", "EDEDED"), "skip": ("Not checked", "EDEDED"),
}
BLOCK_LABELS = [("quality", "Quality", 30), ("growth", "Growth", 25), ("safety", "Safety", 20),
                ("valuation", "Valuation", 15), ("ownership", "Ownership", 10)]
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def sheet_name(symbol, used=()):
    """Excel sheet names: max 31 chars, none of []:*?/\\ and unique (case-insensitive)."""
    base = re.sub(r"[\[\]:*?/\\]", "_", symbol)[:31] or "sheet"
    name, n = base, 2
    taken = {u.lower() for u in used}
    while name.lower() in taken:
        suffix = f"_{n}"
        name, n = base[:31 - len(suffix)] + suffix, n + 1
    return name


def link(cell, location, text=None):
    """Internal (same-workbook) hyperlink."""
    cell.hyperlink = Hyperlink(ref=cell.coordinate, location=location, display=text or str(cell.value))
    cell.font = Font(color="0563C1", underline="single", bold=cell.font.bold, size=cell.font.size)


def _fmt(v, spec, unit=""):
    return format(v, spec) + unit if v is not None else "n/a"


def add_symbol_sheet(wb, symbol, summary, name, tech=None, top=None, watchlist_title="Watchlist"):
    """Create the sheet `name` for one symbol. `summary` is a screen_symbols() entry."""
    ws = wb.create_sheet(name)
    ws.sheet_view.showGridLines = False
    for col, w in zip("ABC", (38, 78, 16)):
        ws.column_dimensions[col].width = w
    top = top or summary.get("top") or {}

    back = ws.cell(row=1, column=1, value=f"<< Back to {watchlist_title}")
    link(back, f"'{watchlist_title}'!A1")

    title = ws.cell(row=2, column=1, value=f"{symbol}" + (f"  -  {summary['industry']}" if summary.get("industry") else ""))
    title.font = Font(bold=True, size=16)
    ws.merge_cells("A2:C2")

    verdict = summary["verdict"]
    ws.cell(row=3, column=1, value="Fundamental verdict").font = Font(bold=True)
    v = ws.cell(row=3, column=2, value=verdict + (f"  -  {summary['score']:g} / 100" if summary.get("score") is not None else ""))
    v.font = Font(bold=True, size=12)
    v.fill = PatternFill("solid", fgColor=VERDICT_FILLS.get(verdict, "FFFFFF"))

    blocks = summary.get("blocks") or {}
    if blocks:
        ws.cell(row=4, column=1, value="Block scores").font = Font(bold=True)
        ws.cell(row=4, column=2, value="   ".join(f"{label} {blocks[k]:g}/{mx}" for k, label, mx in BLOCK_LABELS if k in blocks))

    ws.cell(row=5, column=1, value="Key figures").font = Font(bold=True)
    ws.cell(row=5, column=2, value="   ".join([
        f"Market cap Rs {_fmt(top.get('market_cap'), ',.0f')} Cr", f"Price Rs {_fmt(top.get('price'), ',.2f')}",
        f"P/E {_fmt(top.get('pe'), '.1f')}", f"ROCE {_fmt(top.get('roce'), '.1f')}%", f"ROE {_fmt(top.get('roe'), '.1f')}%",
        f"Book value Rs {_fmt(top.get('book'), ',.0f')}", f"Dividend yield {_fmt(top.get('div_yield'), '.2f')}%"]))

    if tech:
        ws.cell(row=6, column=1, value="Technical (ATH scan)").font = Font(bold=True)
        ws.cell(row=6, column=2, value="   ".join([
            str(tech.get("status", "")), f"Distance {_fmt(tech.get('distance_pct'), '+.2f', '%')}",
            f"Win rate {_fmt(tech.get('win_rate_pct'), '.0f', '%')}", f"Profit factor {_fmt(tech.get('profit_factor'), '.1f')}",
            f"Trades {tech.get('total_trades') if tech.get('total_trades') is not None else 'n/a'}"]))

    row = 8
    head = ws.cell(row=row, column=1, value="Summary")
    head.font = Font(bold=True, color="FFFFFF")
    for c in range(1, 4):
        ws.cell(row=row, column=c).fill = PatternFill("solid", fgColor="404040")
    row += 1
    for line in summary.get("notes") or summary_notes(summary, top, summary.get("basis")):
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=3)
        c = ws.cell(row=row, column=1, value=line)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[row].height = 15 * max(1, -(-len(line) // 125))
        row += 1

    row += 1
    for c, text in enumerate(("Check", "Value", "Result"), start=1):
        cell = ws.cell(row=row, column=c, value=text)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="404040")
        cell.border = BORDER
    row += 1

    results = sorted(summary.get("results", []), key=lambda r: [int(p) for p in r["id"].split(".")])
    last_gate = None
    for r in results:
        if r["gate"] != last_gate:
            last_gate = r["gate"]
            g = ws.cell(row=row, column=1, value=GATE_NAMES.get(r["gate"], f"Gate {r['gate']}"))
            for c in range(1, 4):
                ws.cell(row=row, column=c).fill = PatternFill("solid", fgColor="D9D9D9")
                ws.cell(row=row, column=c).border = BORDER
            g.font = Font(bold=True)
            row += 1
        label, fill = RESULT_STYLE[r["level"]]
        if r["level"] == "fail" and r["gate"] in HARD_GATES:
            label = "Fail - Reject"
        ws.cell(row=row, column=1, value=f"{r['name']} ({r['id']})")
        ws.cell(row=row, column=2, value=r["msg"]).alignment = Alignment(wrap_text=True, vertical="top")
        res = ws.cell(row=row, column=3, value=label)
        res.fill = PatternFill("solid", fgColor=fill)
        res.font = Font(bold=r["level"] in ("warn", "fail"))
        for c in range(1, 4):
            ws.cell(row=row, column=c).border = BORDER
        row += 1
    ws.freeze_panes = "A3"
    return ws
