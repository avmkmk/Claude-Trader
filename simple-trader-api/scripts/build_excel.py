"""
Daily ATH scan routine - Step D.

Merges Step B's candidate data (data/daily_scans/{date}_candidates.json,
used only for market cap tier) with Step C's live TradingView desktop
verdict details (data/daily_scans/{date}_verdicts_details.json), and
writes the final dated Excel watchlist.

Price/ATH/distance-from-ATH come from Step C (the live Pine strategy's
own values), not from tradingview-cli - the two sources define "ATH"
differently (wick high vs. body high), so cross-checking one against
the other produces false rejections. tradingview-cli is only used for
market cap classification.

verdicts_details.json format: list of
    {"symbol": str, "verdict": "reclaimed"|"approaching"|"skip",
     "distance_pct": float, "ema200": float, "absolute_ath": float,
     "close": float, ...}
as written by tradingview-mcp-jackson/scan_step_c.mjs.

Usage:
    python scripts/build_excel.py <candidates.json> <verdicts_details.json>
"""
import json
import os
import sys
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

OUTPUT_DIR = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "data", "daily_scans")

CAP_TIER_ORDER = ["large", "mid", "small", "unknown"]
CAP_TIER_LABELS = {"large": "Large Cap", "mid": "Mid Cap", "small": "Small Cap", "unknown": "Unknown"}

STATUS_LABELS = {"reclaimed": "Reclaimed", "approaching": "Approaching"}


def load_json(path):
    with open(path) as f:
        return json.load(f)


def merge(candidates: list, details: list) -> list:
    cap_by_symbol = {c["symbol"]: c for c in candidates}

    kept = []
    for d in details:
        verdict = d.get("verdict")
        if verdict not in STATUS_LABELS:
            continue

        symbol = d["symbol"]
        cap = cap_by_symbol.get(symbol, {})

        kept.append({
            "symbol": symbol,
            "cap_size": cap.get("cap_size", "unknown"),
            "market_cap_cr": cap.get("market_cap_cr"),
            "status": STATUS_LABELS[verdict],
            "current_price": d.get("close"),
            "all_time_high": d.get("absolute_ath"),
            "ema_200": d.get("ema200"),
            "distance_from_ath_pct": d.get("distance_pct"),
        })

    return kept


def build_excel(kept: list, output_path: str):
    wb = Workbook()
    ws = wb.active
    ws.title = "Watchlist"

    headers = [
        "Symbol", "Cap Tier", "Market Cap (cr)", "Status",
        "Current Price", "ATH", "EMA 200", "Distance from ATH %", "TradingView Link",
    ]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)

    by_tier = {}
    for c in kept:
        by_tier.setdefault(c.get("cap_size", "unknown"), []).append(c)

    row_idx = 2
    for tier in CAP_TIER_ORDER:
        rows = by_tier.get(tier, [])
        if not rows:
            continue
        rows.sort(key=lambda r: r["distance_from_ath_pct"])

        header_cell = ws.cell(row=row_idx, column=1, value=CAP_TIER_LABELS[tier])
        header_cell.font = Font(bold=True)
        header_cell.fill = PatternFill("solid", fgColor="DDDDDD")
        row_idx += 1

        for r in rows:
            ws.cell(row=row_idx, column=1, value=r["symbol"])
            ws.cell(row=row_idx, column=2, value=CAP_TIER_LABELS[tier])
            ws.cell(row=row_idx, column=3, value=r.get("market_cap_cr"))
            ws.cell(row=row_idx, column=4, value=r["status"])
            ws.cell(row=row_idx, column=5, value=r.get("current_price"))
            ws.cell(row=row_idx, column=6, value=r.get("all_time_high"))
            ws.cell(row=row_idx, column=7, value=r.get("ema_200"))
            ws.cell(row=row_idx, column=8, value=r["distance_from_ath_pct"])
            ws.cell(row=row_idx, column=9, value=f"https://in.tradingview.com/chart/?symbol=NSE:{r['symbol']}")
            row_idx += 1

    for col_idx, header in enumerate(headers, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = max(14, len(header) + 2)

    wb.save(output_path)


def main():
    if len(sys.argv) != 3:
        print("Usage: python build_excel.py <candidates.json> <verdicts_details.json>")
        sys.exit(1)

    candidates = load_json(sys.argv[1])
    details = load_json(sys.argv[2])

    kept = merge(candidates, details)

    today = date.today().isoformat()
    output_path = os.path.join(OUTPUT_DIR, f"{today}.xlsx")
    build_excel(kept, output_path)

    print(f"Wrote {len(kept)} stocks to {output_path}")
    print(output_path)
    return output_path


if __name__ == "__main__":
    main()
