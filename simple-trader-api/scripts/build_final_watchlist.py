"""
Builds the dated watchlist Excel from the persistent symbol-state file.

Page 1 ("Watchlist") is ONE clean table - a single header row, no section rows, real numbers and real dates, with
filter/sort dropdowns on every column - so the user can sort and filter it any way they like. The default row order
is Reclaimed first, then Approaching, each by closeness to the entry price (smallest |distance| first); nothing is
locked. Use the Status column to filter to one group.

Reclaimed stocks are filtered/sorted by distance from their ACTUAL entry
price (the RECLAIM bar's close), not the current ATH - the ATH keeps
ratcheting up while a position is held, so "close to ATH" does not mean
"close to entry" (e.g. WELCORP entered at 1005.5 in April, current price
1881 is +87% past entry despite sitting right at the current ATH).
Approaching stocks have no entry yet, so distance-from-ATH is the right
"how close to triggering" measure for them.

Each row also gets fundamental columns from scripts/fundamental_screen.py (docs/FUNDAMENTAL_RULES.md) and a link to
its own sheet. That step fetches screener.in pages (cached per session date, ~3-5 min when fresh); if it fails the
technical watchlist is still written with blank fundamental columns.

Gate 7 (news + government stance) is produced by the Claude command, not by this script: this script writes
{scan_date}_gate7_targets.json (the PASS/WATCH survivors) and, if {scan_date}_gate7.json exists, adds its columns
and a "News and policy" section to each symbol sheet.

Usage:
    python scripts/build_final_watchlist.py <scan_date> <output_path> [--no-fundamentals] [--refresh-fundamentals]
"""
import json
import os
import sys
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

DATA_DIR = "data/daily_scans"
STATE_PATH = f"{DATA_DIR}/symbol_state.json"
EXCLUSIONS_PATH = f"{DATA_DIR}/manual_exclusions.json"

FUNDAMENTAL_HEADERS = [
    "Fundamental Verdict", "Fundamental Score", "Quality /30", "Growth /25", "Safety /20",
    "Valuation /15", "Ownership /10", "Industry", "Key Note",
]
GATE7_HEADERS = ["News Sentiment", "Policy Stance", "News & Policy Note"]
VERDICT_FILLS = {"PASS": "C6EFCE", "WATCH": "FFEB9C", "REJECT": "FFC7CE", "FINANCIAL": "DDEBF7", "NO DATA": "D9D9D9"}
SENTIMENT_FILLS = {"Bullish": "C6EFCE", "Neutral": "EDEDED", "Bearish": "FFC7CE"}
SENTIMENTS = ("Bullish", "Neutral", "Bearish")
SURVIVOR_VERDICTS = ("PASS", "WATCH")

ENTRY_BAND_PCT = 5.0
CAP_LABELS = {"large": "Large Cap", "mid": "Mid Cap", "small": "Small Cap", "unknown": "Unknown"}


def load(path):
    with open(path, encoding="utf-8") as f:
        return json.load(f)


def load_exclusions():
    try:
        return load(EXCLUSIONS_PATH)
    except FileNotFoundError:
        return {}


def compute_combined_rank(rows):
    eligible = [r for r in rows if r.get("ranking_eligible")]

    def dense_rank_by(key):
        # Ties share a rank (distinct sorted values, not list position) so
        # two symbols with identical metrics get an identical combined_rank.
        distinct_values = sorted({r[key] for r in eligible}, reverse=True)
        return {v: i for i, v in enumerate(distinct_values)}

    wr_rank = dense_rank_by("win_rate_pct")
    pf_rank = dense_rank_by("profit_factor")
    for r in rows:
        if r.get("ranking_eligible"):
            r["combined_rank"] = (wr_rank[r["win_rate_pct"]] + pf_rank[r["profit_factor"]]) / 2
        else:
            r["combined_rank"] = None


def key_note(summary):
    """One short line for page 1: the first hard fail, else the first warning (full detail is on the symbol's sheet)."""
    if summary["verdict"] == "FINANCIAL":
        return "Bank / NBFC / insurer - not scored"
    for text in list(summary.get("hard_fails") or []) + list(summary.get("flags") or []):
        text = text.split(": ", 1)[1] if ": " in text else text
        return text if len(text) <= 90 else text[:87] + "..."
    return None


def fundamental_cells(summary):
    """Values for FUNDAMENTAL_HEADERS from a screen_symbols() summary (empty cells if None)."""
    if not summary:
        return [None] * len(FUNDAMENTAL_HEADERS)
    b = summary.get("blocks") or {}
    return [
        summary["verdict"], summary["score"], b.get("quality"), b.get("growth"), b.get("safety"),
        b.get("valuation"), b.get("ownership"), summary.get("industry"), key_note(summary),
    ]


def clean_gate7(entry):
    """Normalise one Gate 7 entry written by the Claude command."""
    def sent(v):
        v = str(v or "").strip().capitalize()
        return v if v in SENTIMENTS else "Unrated"
    return {
        "news_sentiment": sent(entry.get("news_sentiment")), "policy_sentiment": sent(entry.get("policy_sentiment")),
        "summary": str(entry.get("summary") or "").strip(),
        "headlines": [h for h in (entry.get("headlines") or []) if isinstance(h, dict) and h.get("title")],
        "checked_on": entry.get("checked_on"),
    }


def load_gate7(scan_date):
    """{symbol: cleaned entry} from {scan_date}_gate7.json, or {} if the Claude command has not run yet."""
    path = f"{DATA_DIR}/{scan_date}_gate7.json"
    if not os.path.exists(path):
        return {}
    return {sym: clean_gate7(e) for sym, e in load(path).items() if isinstance(e, dict)}


def gate7_cells(symbol, summary, gate7, gate7_ran):
    """Values for GATE7_HEADERS: real data, 'Pending' (survivor, not run yet) or 'n/a' (Gate 7 only covers PASS/WATCH)."""
    g = gate7.get(symbol)
    if g:
        note = g["summary"] if len(g["summary"]) <= 160 else g["summary"][:157] + "..."
        return [g["news_sentiment"], g["policy_sentiment"], note]
    if summary and summary["verdict"] in SURVIVOR_VERDICTS:
        return ["Pending" if not gate7_ran else "Not covered"] * 2 + [None]
    return ["n/a", "n/a", None]


def parse_date(text):
    try:
        return date.fromisoformat(text) if text else None
    except (TypeError, ValueError):
        return None


def build_rows(state, exclusions):
    """Watchlist rows (dicts) in default order: Reclaimed then Approaching, each by |distance| ascending."""
    reclaimed, approaching = [], []
    for symbol, r in state.items():
        if symbol in exclusions:
            continue
        cap_size = r.get("cap_size", "unknown")
        market_cap_cr = r.get("market_cap_cr")
        if r.get("verdict") == "reclaimed":
            dist = r.get("distance_from_entry_pct")
            if dist is None or abs(dist) > ENTRY_BAND_PCT:
                continue
            reclaimed.append({
                "symbol": symbol, "status": "Reclaimed", "cap_size": cap_size,
                "market_cap_cr": market_cap_cr, "distance_pct": dist,
                "close": r.get("close"), "reference_price": r.get("entry_price"),
                "reference_label": "Entry Price", "ema200": r.get("ema200"),
                "win_rate_pct": r.get("win_rate_pct"), "profit_factor": r.get("profit_factor"),
                "total_trades": r.get("total_trades"), "ranking_eligible": r.get("ranking_eligible", False),
                "entry_date": r.get("entry_date"), "last_checked": r.get("last_checked"),
            })
        elif r.get("verdict") == "approaching":
            approaching.append({
                "symbol": symbol, "status": "Approaching", "cap_size": cap_size,
                "market_cap_cr": market_cap_cr, "distance_pct": r.get("distance_pct"),
                "close": r.get("close"), "reference_price": r.get("absolute_ath"),
                "reference_label": "ATH (entry trigger)", "ema200": r.get("ema200"),
                "win_rate_pct": r.get("win_rate_pct"), "profit_factor": r.get("profit_factor"),
                "total_trades": r.get("total_trades"), "ranking_eligible": r.get("ranking_eligible", False),
                "entry_date": None, "last_checked": r.get("last_checked"),
            })
    compute_combined_rank(reclaimed)
    compute_combined_rank(approaching)
    key = lambda r: abs(r["distance_pct"]) if r["distance_pct"] is not None else 1e9  # noqa: E731
    return sorted(reclaimed, key=key) + sorted(approaching, key=key), len(reclaimed), len(approaching)


def write_workbook(output_path, rows, fundamentals, gate7, gate7_ran=False):
    """One filterable Watchlist table plus one sheet per symbol (when fundamentals are available)."""
    sheet_names = {}
    if fundamentals:
        from scripts.fundamental_sheets import add_symbol_sheet, link, sheet_name
        used = ["Watchlist"]
        for r in rows:
            if r["symbol"] in fundamentals:
                sheet_names[r["symbol"]] = sheet_name(r["symbol"], used)
                used.append(sheet_names[r["symbol"]])

    wb = Workbook()
    ws = wb.active
    ws.title = "Watchlist"
    headers = [
        "Symbol", "Status", "Cap Tier", "Market Cap (cr)", "Distance from Entry %",
        "Current Price", "Reference Price", "Reference", "EMA 200",
        "Win Rate %", "Profit Factor", "Total Trades", "Combined Rank",
        *FUNDAMENTAL_HEADERS, *GATE7_HEADERS,
        "Entry Date", "Last Checked", "TradingView Link",
    ]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    col_of = {h: i for i, h in enumerate(headers, start=1)}

    for r in rows:
        summ = fundamentals.get(r["symbol"])
        values = [
            r["symbol"], r["status"], CAP_LABELS.get(r["cap_size"], "Unknown"), r["market_cap_cr"],
            r["distance_pct"], r["close"], r["reference_price"], r["reference_label"], r["ema200"],
            r["win_rate_pct"], r["profit_factor"], r["total_trades"], r["combined_rank"],
            *fundamental_cells(summ), *gate7_cells(r["symbol"], summ, gate7, gate7_ran),
            parse_date(r["entry_date"]), parse_date(r["last_checked"]),
            f"https://in.tradingview.com/chart/?symbol=NSE:{r['symbol']}",
        ]
        ws.append(values)
        row = ws.max_row
        for h in ("Entry Date", "Last Checked"):
            ws.cell(row=row, column=col_of[h]).number_format = "yyyy-mm-dd"
        if r["symbol"] in sheet_names:  # the symbol cell jumps to that stock's sheet
            link(ws.cell(row=row, column=1), f"'{sheet_names[r['symbol']]}'!A1")
        verdict = ws.cell(row=row, column=col_of["Fundamental Verdict"]).value
        if verdict in VERDICT_FILLS:
            ws.cell(row=row, column=col_of["Fundamental Verdict"]).fill = PatternFill("solid", fgColor=VERDICT_FILLS[verdict])
        for h in ("News Sentiment", "Policy Stance"):
            v = ws.cell(row=row, column=col_of[h]).value
            if v in SENTIMENT_FILLS:
                ws.cell(row=row, column=col_of[h]).fill = PatternFill("solid", fgColor=SENTIMENT_FILLS[v])

    wide = {"Key Note": 60, "News & Policy Note": 60, "Industry": 24}
    for idx, header in enumerate(headers, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = wide.get(header, max(14, len(header) + 2))
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{ws.max_row}"  # dropdown sort/filter on every column

    for r in rows:  # one sheet per symbol, in watchlist order
        if r["symbol"] in sheet_names:
            add_symbol_sheet(wb, r["symbol"], fundamentals[r["symbol"]], sheet_names[r["symbol"]], tech=r,
                             gate7=gate7.get(r["symbol"]), gate7_ran=gate7_ran)
    wb.save(output_path)
    return sheet_names


def main():
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    flags = {a for a in sys.argv[1:] if a.startswith("--")}
    if len(args) != 2 or not flags <= {"--no-fundamentals", "--refresh-fundamentals"}:
        print("Usage: python build_final_watchlist.py <scan_date> <output_path> [--no-fundamentals] [--refresh-fundamentals]")
        sys.exit(1)
    scan_date, output_path = args

    rows, n_reclaimed, n_approaching = build_rows(load(STATE_PATH), load_exclusions())

    fundamentals = {}
    if "--no-fundamentals" not in flags and rows:
        try:
            sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))  # so `scripts.*` imports work when run as a script
            from scripts.fundamental_screen import screen_symbols
            print(f"Running fundamental screen on {len(rows)} symbols (cached per session date)...")
            fundamentals = screen_symbols([r["symbol"] for r in rows], refresh="--refresh-fundamentals" in flags, cache_date=scan_date)
        except Exception as e:  # never lose the technical watchlist over the fundamental step
            print(f"WARNING: fundamental screen failed ({type(e).__name__}: {e}); writing watchlist without fundamentals")

    gate7 = load_gate7(scan_date)
    write_workbook(output_path, rows, fundamentals, gate7, gate7_ran=bool(gate7))

    # Hand the PASS/WATCH survivors to the Claude command for Gate 7 (news + government stance).
    targets = [{"symbol": r["symbol"], "name": fundamentals[r["symbol"]].get("name") or r["symbol"],
                "industry": fundamentals[r["symbol"]].get("industry"), "status": r["status"],
                "verdict": fundamentals[r["symbol"]]["verdict"], "score": fundamentals[r["symbol"]]["score"],
                "market_cap_cr": r["market_cap_cr"]}
               for r in rows if r["symbol"] in fundamentals and fundamentals[r["symbol"]]["verdict"] in SURVIVOR_VERDICTS]
    with open(f"{DATA_DIR}/{scan_date}_gate7_targets.json", "w", encoding="utf-8") as f:
        json.dump(targets, f, indent=2)

    print(f"Reclaimed near entry: {n_reclaimed}")
    print(f"Approaching: {n_approaching}")
    if fundamentals:
        tally = {}
        for summ in fundamentals.values():
            tally[summ["verdict"]] = tally.get(summ["verdict"], 0) + 1
        print("Fundamentals: " + ", ".join(f"{k}={v}" for k, v in sorted(tally.items())))
        print(f"Gate 7 targets (PASS/WATCH): {len(targets)}" + (f"; Gate 7 results merged for {len(gate7)}" if gate7 else ""))
    print(f"Wrote {len(rows)} stocks to {output_path}")


if __name__ == "__main__":
    main()
