"""
Run the fundamental rules (docs/FUNDAMENTAL_RULES.md) over the latest daily ATH watchlist.

Usage (from simple-trader-api/):
    python scripts/fundamental_screen.py [watchlist.xlsx] [--delay 2.0] [--refresh]

Reads the newest data/daily_scans/*_final_watchlist.xlsx (or the one given), fetches each symbol's
screener.in page (consolidated, falling back to standalone), caches the parsed data per day under
data/fundamentals/{date}/, and writes data/daily_scans/{date}_fundamental_screen.xlsx.
"""
import argparse
import glob
import json
import os
import re
import sys
import time
from datetime import date

import requests
from bs4 import BeautifulSoup
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from scripts.fundamental_rules import evaluate  # noqa: E402

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SCANS_DIR = os.path.join(BASE, "data", "daily_scans")
CACHE_DIR = os.path.join(BASE, "data", "fundamentals")
HEADERS = {"User-Agent": "Mozilla/5.0 (personal research script)"}
TABLES = ["quarters", "profit-loss", "balance-sheet", "cash-flow", "ratios", "shareholding"]


def to_num(text):
    t = (text or "").replace(",", "").replace("%", "").replace("₹", "").replace("Cr.", "").strip()
    if t in ("", "-"):
        return None
    try:
        return float(t)
    except ValueError:
        return None


def parse_table(section):
    table = section.find("table") if section else None
    if table is None:
        return None
    trs = table.find_all("tr")
    periods = [th.get_text(strip=True) for th in trs[0].find_all("th")][1:]
    rows = {}
    for tr in trs[1:]:
        cells = tr.find_all("td")
        if len(cells) < 2:
            continue
        name = re.sub(r"\s+", " ", cells[0].get_text(" ", strip=True)).replace("+", "").strip()
        rows[name] = [to_num(c.get_text(strip=True)) for c in cells[1:]]
    return {"periods": periods, "rows": rows}


def parse_peers_median_pe(html):
    soup = BeautifulSoup(html, "html.parser")
    trs = soup.find_all("tr")
    if not trs:
        return None
    heads = [c.get_text(" ", strip=True) for c in trs[0].find_all(["th", "td"])]
    if "P/E" not in heads:
        return None
    idx = heads.index("P/E")
    for tr in trs:
        cells = [c.get_text(" ", strip=True) for c in tr.find_all(["th", "td"])]
        if len(cells) > idx and cells[1:2] and cells[1].lower().startswith("median"):
            return to_num(cells[idx])
    return None


def parse_company_page(html, peers_html=None):
    """screener.in company page -> dict consumed by fundamental_rules.evaluate."""
    soup = BeautifulSoup(html, "html.parser")
    top = {}
    keymap = {"Market Cap": "market_cap", "Current Price": "price", "Stock P/E": "pe", "Book Value": "book",
              "Dividend Yield": "div_yield", "ROCE": "roce", "ROE": "roe"}
    for li in soup.select("#top-ratios li"):
        name = li.find(class_="name")
        val = li.find(class_="number")
        if name and val and name.get_text(strip=True) in keymap:
            top[keymap[name.get_text(strip=True)]] = to_num(val.get_text(strip=True))
    sector = {}
    peers = soup.find("section", id="peers")
    if peers:
        titles = {"Broad Sector": "broad", "Sector": "sector", "Broad Industry": "broad_industry", "Industry": "industry"}
        for a in peers.select("p.sub a"):
            if a.get("title") in titles:
                sector[titles[a.get("title")]] = a.get_text(strip=True)
    tables = {}
    for tid in TABLES:
        t = parse_table(soup.find("section", id=tid))
        if t:
            tables[tid] = t
    growth = {}
    for t in soup.select("table.ranges-table"):
        title = t.find("th").get_text(strip=True)
        growth[title] = {}
        for tr in t.find_all("tr")[1:]:
            if ":" in tr.get_text():
                label, val = tr.get_text(" ", strip=True).split(":", 1)
                growth[title][label.strip()] = to_num(val)
    wid = re.search(r'data-warehouse-id="(\d+)"', html)
    return {"top": top, "sector": sector, "tables": tables, "growth": growth,
            "peers_median_pe": parse_peers_median_pe(peers_html) if peers_html else None,
            "warehouse_id": wid.group(1) if wid else None}


def http_get(url, extra=None, retries=3):
    for attempt in range(retries):
        r = requests.get(url, headers={**HEADERS, **(extra or {})}, timeout=30)
        if r.status_code == 429:
            time.sleep(10 * (attempt + 1))
            continue
        return r
    return r


def _usable(data):
    """(annual years, quarters) with real values - how much history a parsed page actually has."""
    pl = data["tables"].get("profit-loss", {})
    years = len([v for v in pl.get("rows", {}).get("Sales", []) if v is not None])
    qtrs = len([v for v in data["tables"].get("quarters", {}).get("rows", {}).get("Sales", []) if v is not None])
    return years, qtrs


def fetch_symbol(symbol, delay):
    """Return (parsed data, None) or (None, reason).

    Tries consolidated first. If its history is thin (< 5 years or < 5 quarters - empty or recently
    consolidated pages), also tries standalone and keeps whichever has more annual years.
    """
    candidates, last_err = [], "no financial tables"
    for suffix in ("/consolidated/", "/"):
        r = http_get(f"https://www.screener.in/company/{symbol}{suffix}")
        time.sleep(delay)
        if r.status_code == 404:
            last_err = "not found on screener.in"
            break
        if r.status_code != 200:
            last_err = f"HTTP {r.status_code}"
            break
        if 'id="profit-loss"' not in r.text:
            continue
        data = parse_company_page(r.text)
        data["symbol"], data["basis"] = symbol, "consolidated" if "consolidated" in suffix else "standalone"
        data["_html"] = r.text
        candidates.append(data)  # keep even if empty: banks/NBFCs have no "Sales" row but are routed by sector
        years, qtrs = _usable(data)
        if suffix == "/consolidated/" and years >= 5 and qtrs >= 5:
            break
    if not candidates:
        return None, last_err
    best = max(candidates, key=lambda c: (_usable(c)[0], c["basis"] == "consolidated"))
    html = best.pop("_html")
    for c in candidates:
        c.pop("_html", None)
    if best.get("warehouse_id"):
        pr = http_get(f"https://www.screener.in/api/company/{best['warehouse_id']}/peers/", {"X-Requested-With": "XMLHttpRequest"})
        time.sleep(delay)
        if pr.status_code == 200:
            best["peers_median_pe"] = parse_peers_median_pe(pr.text)
    return best, None


def latest_watchlist():
    files = sorted(glob.glob(os.path.join(SCANS_DIR, "*_final_watchlist.xlsx")))
    if not files:
        sys.exit("No *_final_watchlist.xlsx found in data/daily_scans/")
    return files[-1]


def read_watchlist(path):
    ws = load_workbook(path, data_only=True).active
    header = [c.value for c in ws[1]]
    out = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        rec = dict(zip(header, row))
        if rec.get("Symbol") and rec.get("Status"):
            out.append(rec)
    return header, out


FILLS = {"PASS": "C6EFCE", "WATCH": "FFEB9C", "REJECT": "FFC7CE", "FINANCIAL": "DDEBF7", "NO DATA": "D9D9D9"}


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("watchlist", nargs="?")
    ap.add_argument("--delay", type=float, default=2.0)
    ap.add_argument("--refresh", action="store_true", help="ignore today's cache")
    args = ap.parse_args()

    path = args.watchlist or latest_watchlist()
    scan_date = re.match(r"(\d{4}-\d{2}-\d{2})", os.path.basename(path))
    scan_date = scan_date.group(1) if scan_date else date.today().isoformat()
    cache_dir = os.path.join(CACHE_DIR, date.today().isoformat())
    os.makedirs(cache_dir, exist_ok=True)

    header, rows = read_watchlist(path)
    print(f"Watchlist: {os.path.basename(path)} ({len(rows)} symbols)")

    out_rows, detail_rows = [], []
    for i, rec in enumerate(rows, 1):
        sym = rec["Symbol"]
        cache = os.path.join(cache_dir, f"{sym}.json")
        data, err = None, None
        if os.path.exists(cache) and not args.refresh:
            data = json.load(open(cache, encoding="utf-8"))
        else:
            try:
                data, err = fetch_symbol(sym, args.delay)
            except requests.RequestException as e:
                err = f"network error: {e}"
            if data:
                json.dump(data, open(cache, "w", encoding="utf-8"))
        if data is None:
            res = {"verdict": "NO DATA", "score": None, "blocks": {}, "results": [], "hard_fails": [err]}
        else:
            res = evaluate(data)
        warns = [f'{r["id"]}: {r["msg"]}' for r in res["results"] if r["level"] == "warn"]
        fails = [f'{r["id"]}: {r["msg"]}' for r in res["results"] if r["level"] == "fail"]
        out_rows.append({"rec": rec, "res": res, "sector": (data or {}).get("sector", {}), "flags": fails + warns})
        for r in res["results"]:
            detail_rows.append([sym, r["id"], r["level"].upper(), r["msg"]])
        print(f"[{i}/{len(rows)}] {sym:12s} {res['verdict']:9s} {res['score'] if res['score'] is not None else '':>5}  {'; '.join(res['hard_fails'])[:90]}")

    order = {"PASS": 0, "WATCH": 1, "REJECT": 2, "FINANCIAL": 3, "NO DATA": 4}
    out_rows.sort(key=lambda x: (order[x["res"]["verdict"]], -(x["res"]["score"] or 0)))

    wb = Workbook()
    ws = wb.active
    ws.title = "Fundamentals"
    cols = ["Symbol", "Status", "Cap Tier", "Market Cap (cr)", "Industry", "Fundamental Verdict", "Score", "Quality /30",
            "Growth /25", "Safety /20", "Valuation /15", "Ownership /10", "Hard Fails", "Warnings", "TradingView Link"]
    ws.append(cols)
    for c in ws[1]:
        c.font = Font(bold=True)
    for x in out_rows:
        rec, res, b = x["rec"], x["res"], x["res"]["blocks"]
        ws.append([rec["Symbol"], rec["Status"], rec.get("Cap Tier"), rec.get("Market Cap (cr)"), x["sector"].get("industry"),
                   res["verdict"], res["score"], b.get("quality"), b.get("growth"), b.get("safety"), b.get("valuation"),
                   b.get("ownership"), "\n".join(res["hard_fails"]), "\n".join(x["flags"]), rec.get("TradingView Link")])
        ws.cell(ws.max_row, 6).fill = PatternFill("solid", fgColor=FILLS[res["verdict"]])
    for i, w in enumerate([12, 11, 11, 14, 22, 18, 8, 10, 10, 10, 11, 11, 60, 70, 45], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "B2"
    d = wb.create_sheet("Rule Details")
    d.append(["Symbol", "Rule", "Result", "Detail"])
    for c in d[1]:
        c.font = Font(bold=True)
    for r in sorted(detail_rows, key=lambda r: (r[0], [int(p) for p in r[1].split(".")])):
        d.append(r)
    for i, w in enumerate([12, 7, 9, 90], 1):
        d.column_dimensions[get_column_letter(i)].width = w
    out = os.path.join(SCANS_DIR, f"{scan_date}_fundamental_screen.xlsx")
    wb.save(out)

    tally = {}
    for x in out_rows:
        tally[x["res"]["verdict"]] = tally.get(x["res"]["verdict"], 0) + 1
    print("\nSummary:", ", ".join(f"{k}={v}" for k, v in sorted(tally.items(), key=lambda kv: order[kv[0]])))
    print("Wrote", out)


if __name__ == "__main__":
    main()
