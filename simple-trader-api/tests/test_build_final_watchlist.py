"""Tests for build_final_watchlist.py's ranking logic."""
from scripts.build_final_watchlist import compute_combined_rank


def test_ranks_eligible_symbols_by_win_rate_and_profit_factor():
    rows = [
        {"symbol": "A", "ranking_eligible": True, "win_rate_pct": 80.0, "profit_factor": 2.0},
        {"symbol": "B", "ranking_eligible": True, "win_rate_pct": 60.0, "profit_factor": 10.0},
        {"symbol": "C", "ranking_eligible": True, "win_rate_pct": 90.0, "profit_factor": 8.0},
    ]
    compute_combined_rank(rows)
    ranks = {r["symbol"]: r["combined_rank"] for r in rows}
    # C: best win_rate (rank 0) + 2nd best profit_factor (rank 1) -> 0.5
    # A: 2nd best win_rate (rank 1) + worst profit_factor (rank 2) -> 1.5
    # B: worst win_rate (rank 2) + best profit_factor (rank 0) -> 1.0
    assert ranks["C"] < ranks["B"] < ranks["A"]


def test_ineligible_symbols_get_null_rank():
    rows = [
        {"symbol": "A", "ranking_eligible": True, "win_rate_pct": 80.0, "profit_factor": 2.0},
        {"symbol": "B", "ranking_eligible": False, "win_rate_pct": None, "profit_factor": None},
    ]
    compute_combined_rank(rows)
    assert rows[0]["combined_rank"] is not None
    assert rows[1]["combined_rank"] is None


def test_empty_list_does_not_error():
    rows = []
    compute_combined_rank(rows)
    assert rows == []


def test_all_ineligible_all_null():
    rows = [
        {"symbol": "A", "ranking_eligible": False, "win_rate_pct": None, "profit_factor": None},
        {"symbol": "B", "ranking_eligible": False, "win_rate_pct": None, "profit_factor": None},
    ]
    compute_combined_rank(rows)
    assert all(r["combined_rank"] is None for r in rows)


def test_tied_metrics_produce_equal_rank():
    rows = [
        {"symbol": "A", "ranking_eligible": True, "win_rate_pct": 70.0, "profit_factor": 3.0},
        {"symbol": "B", "ranking_eligible": True, "win_rate_pct": 70.0, "profit_factor": 3.0},
    ]
    compute_combined_rank(rows)
    assert rows[0]["combined_rank"] == rows[1]["combined_rank"]


def test_fundamental_cells_blank_when_missing():
    from scripts.build_final_watchlist import fundamental_cells, FUNDAMENTAL_HEADERS
    cells = fundamental_cells(None)
    assert len(cells) == len(FUNDAMENTAL_HEADERS) and set(cells) == {None}


def test_fundamental_cells_maps_summary():
    from scripts.build_final_watchlist import fundamental_cells, FUNDAMENTAL_HEADERS
    summary = {"verdict": "WATCH", "score": 65.4, "blocks": {"quality": 22.1, "growth": 18.6, "safety": 15.2,
               "valuation": 4.5, "ownership": 5.0}, "industry": "Paints", "hard_fails": [],
               "flags": ["1.3: Borrowings up", "3.3: weak"]}
    cells = fundamental_cells(summary)
    assert len(cells) == len(FUNDAMENTAL_HEADERS)
    assert cells[:7] == ["WATCH", 65.4, 22.1, 18.6, 15.2, 4.5, 5.0]
    assert cells[7] == "Paints" and cells[8] == "Borrowings up"


def test_key_note_prefers_hard_fail_and_truncates():
    from scripts.build_final_watchlist import key_note
    long = "0.2: " + "x" * 200
    assert key_note({"verdict": "REJECT", "hard_fails": [long], "flags": ["1.3: other"]}) == "x" * 87 + "..."
    assert key_note({"verdict": "PASS", "hard_fails": [], "flags": []}) is None
    assert key_note({"verdict": "FINANCIAL", "hard_fails": [], "flags": []}).startswith("Bank")


# ---------- workbook layout: one clean, sortable table ----------
from datetime import date  # noqa: E402

from openpyxl import load_workbook  # noqa: E402

from scripts.build_final_watchlist import (GATE7_HEADERS, build_rows, clean_gate7, gate7_cells,  # noqa: E402
                                           write_workbook)
from scripts.fundamental_rules import evaluate, summary_notes  # noqa: E402
from test_fundamental_rules import good_company  # noqa: E402

STATE = {
    "AAA": {"verdict": "reclaimed", "distance_from_entry_pct": 3.0, "entry_price": 100, "close": 103, "cap_size": "large",
            "market_cap_cr": 20000, "entry_date": "2026-05-07", "last_checked": "2026-10-01", "total_trades": 5},
    "BBB": {"verdict": "reclaimed", "distance_from_entry_pct": -0.5, "entry_price": 50, "close": 49.7, "cap_size": "mid",
            "market_cap_cr": 8000, "entry_date": "2026-09-01", "last_checked": "2026-10-01"},
    "CCC": {"verdict": "approaching", "distance_pct": -1.0, "absolute_ath": 200, "close": 198, "cap_size": "small",
            "market_cap_cr": 2000, "last_checked": "2026-10-01"},
    "DDD": {"verdict": "reclaimed", "distance_from_entry_pct": 40.0, "cap_size": "large"},   # too far past entry
    "EEE": {"verdict": "skip"},
}


def fake_summary(verdict="PASS"):
    d = good_company()
    res = evaluate(d)
    summ = {**res, "flags": [], "industry": "Machinery", "basis": "consolidated", "top": d["top"], "name": "Test Co Ltd"}
    summ["verdict"] = verdict
    summ["notes"] = summary_notes(summ, d["top"], "consolidated")
    return summ


def test_build_rows_default_order_and_filtering():
    rows, n_rec, n_app = build_rows(STATE, {})
    assert [r["symbol"] for r in rows] == ["BBB", "AAA", "CCC"] and (n_rec, n_app) == (2, 1)
    assert [r["symbol"] for r in build_rows(STATE, {"BBB": {}})[0]] == ["AAA", "CCC"]


def test_watchlist_is_one_clean_sortable_table(tmp_path):
    rows, _, _ = build_rows(STATE, {})
    out = tmp_path / "w.xlsx"
    write_workbook(str(out), rows, {r["symbol"]: fake_summary() for r in rows}, {})
    ws = load_workbook(out)["Watchlist"]
    assert ws.auto_filter.ref == f"A1:{ws.cell(row=1, column=ws.max_column).column_letter}{ws.max_row}"
    headers = [c.value for c in ws[1]]
    body = [dict(zip(headers, [c.value for c in r])) for r in ws.iter_rows(min_row=2)]
    assert [b["Symbol"] for b in body] == ["BBB", "AAA", "CCC"]          # no "Reclaimed"/"Approaching" separator rows
    assert all(b["Status"] for b in body) and all(b["Fundamental Verdict"] for b in body)
    assert body[1]["Entry Date"].date() == date(2026, 5, 7) and body[2]["Entry Date"] is None   # real dates / real blanks
    assert isinstance(body[0]["Distance from Entry %"], float) and isinstance(body[0]["Fundamental Score"], (int, float))
    assert all(c.hyperlink and c.hyperlink.location == f"'{c.value}'!A1" for c in ws["A"][1:])


def test_blank_cells_stay_empty_without_fundamentals(tmp_path):
    rows, _, _ = build_rows(STATE, {})
    out = tmp_path / "w.xlsx"
    write_workbook(str(out), rows, {}, {})
    ws = load_workbook(out)["Watchlist"]
    headers = [c.value for c in ws[1]]
    assert ws.cell(row=2, column=headers.index("Fundamental Verdict") + 1).value is None
    assert load_workbook(out).sheetnames == ["Watchlist"]


def test_gate7_columns_pending_na_and_merged(tmp_path):
    summ_pass, summ_rej = fake_summary("PASS"), fake_summary("REJECT")
    assert gate7_cells("X", summ_pass, {}, False) == ["Pending", "Pending", None]
    assert gate7_cells("X", summ_rej, {}, False) == ["n/a", "n/a", None]
    g = {"AAA": clean_gate7({"news_sentiment": "bullish", "policy_sentiment": "WEIRD", "summary": "Order win.",
                             "headlines": [{"title": "Big order", "source": "ET", "date": "2026-10-01", "url": "https://x.test/a"}, {"nope": 1}],
                             "checked_on": "2026-10-05"})}
    assert g["AAA"]["news_sentiment"] == "Bullish" and g["AAA"]["policy_sentiment"] == "Unrated" and len(g["AAA"]["headlines"]) == 1
    rows, _, _ = build_rows(STATE, {})
    out = tmp_path / "w.xlsx"
    write_workbook(str(out), rows, {r["symbol"]: fake_summary() for r in rows}, g, gate7_ran=True)
    wb = load_workbook(out)
    ws = wb["Watchlist"]
    headers = [c.value for c in ws[1]]
    assert all(h in headers for h in GATE7_HEADERS)
    by_sym = {r[0].value: dict(zip(headers, [c.value for c in r])) for r in ws.iter_rows(min_row=2)}
    assert by_sym["AAA"]["News Sentiment"] == "Bullish" and by_sym["BBB"]["News Sentiment"] == "Not covered"
    text = [c.value for row in wb["AAA"].iter_rows() for c in row if c.value]
    assert any("Big order" in str(t) for t in text) and any("Order win." == str(t) for t in text)


def test_symbol_sheet_rules_table_is_filterable(tmp_path):
    rows, _, _ = build_rows(STATE, {})
    out = tmp_path / "w.xlsx"
    write_workbook(str(out), rows, {r["symbol"]: fake_summary() for r in rows}, {})
    sheet = load_workbook(out)["AAA"]
    ref = sheet.auto_filter.ref
    assert ref and ref.startswith("A") and ref.split(":")[0][1:].isdigit()
    first = int(ref.split(":")[0][1:])
    assert [c.value for c in sheet[first]] == ["Check", "Value", "Result", "Gate"]
    gates = {sheet.cell(row=r, column=4).value for r in range(first + 1, sheet.max_row + 1)}
    assert None not in gates and any(g.startswith("1 ") for g in gates)   # every row carries a Gate value; no sub-header rows
    assert not sheet.merged_cells.ranges or all(m.min_row < first for m in sheet.merged_cells.ranges)  # nothing merged inside the table
