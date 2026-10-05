"""Tests for the per-symbol sheets and the watchlist hyperlinks."""
from openpyxl import Workbook

from scripts.fundamental_rules import summary_notes
from scripts.fundamental_sheets import add_symbol_sheet, sheet_name
from test_fundamental_rules import good_company
from scripts.fundamental_rules import evaluate


def make_summary():
    d = good_company()
    res = evaluate(d)
    summ = {**res, "flags": [], "industry": "Machinery", "basis": "consolidated", "top": d["top"]}
    summ["notes"] = summary_notes(summ, d["top"], "consolidated")
    return summ


def test_sheet_name_rules():
    assert sheet_name("M&M") == "M&M"
    assert sheet_name("A/B[1]:x") == "A_B_1__x"
    assert len(sheet_name("X" * 50)) == 31
    assert sheet_name("abc", used=["ABC"]) == "abc_2"


def test_symbol_sheet_has_banner_notes_and_every_rule():
    wb = Workbook()
    wb.active.title = "Watchlist"
    summ = make_summary()
    ws = add_symbol_sheet(wb, "GOOD", summ, "GOOD", tech={"status": "Reclaimed", "distance_pct": -0.5, "win_rate_pct": 80, "profit_factor": 2.0, "total_trades": 5})
    text = [c.value for row in ws.iter_rows() for c in row if c.value is not None]
    assert any("PASS" in str(t) and "/ 100" in str(t) for t in text)
    assert any(str(t).startswith("<< Back to Watchlist") for t in text)
    assert ws["A1"].hyperlink.location == "'Watchlist'!A1"
    for r in summ["results"]:
        assert any(f"({r['id']})" in str(t) for t in text), r["id"]
    assert len(summ["notes"]) == 4


def test_notes_for_financial_and_no_data():
    assert "not scored" in summary_notes({"verdict": "FINANCIAL", "results": []})[0]
    assert "No fundamental data" in summary_notes({"verdict": "NO DATA", "hard_fails": ["HTTP 500"], "results": []})[0]
